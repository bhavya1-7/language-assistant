import io
import json
import streamlit as st
import pandas as pd
from docx import Document
from utils import load_data, call_ollama

st.set_page_config(page_title="Quiz & Reports", layout="wide")
students_df, lessons_df = load_data()

st.title("📝 Automated Quizzes & Report Generation")
st.caption("Auto-grading assessments and one-click Word/Excel report exports")

tab1, tab2 = st.tabs(["Self-Checking Quiz", "Report Exports"])

with tab1:
    st.subheader("Adaptive MCQ Assessment")
    lvl = st.selectbox("Select Target Difficulty", ["Beginner", "Intermediate", "Advanced"])
    
    if "current_quiz" not in st.session_state:
        st.session_state.current_quiz = [
            {"q": "Choose the correct verb: The team of developers ___ finishing the release.", "opts": ["A) is", "B) are", "C) were", "D) have"], "ans": "A", "exp": "'The team' is a collective singular subject here."},
            {"q": "Select the closest synonym for 'Comprehensive':", "opts": ["A) Limited", "B) Exhaustive", "C) Superficial", "D) Fragile"], "ans": "B", "exp": "Exhaustive means complete and thorough."},
            {"q": "Identify the passive voice sentence:", "opts": ["A) She wrote the report.", "B) The code was compiled by the build server.", "C) We attended the meeting.", "D) He will submit it."], "ans": "B", "exp": "The subject receives the action via 'was compiled by'."}
        ]
        
    user_answers = {}
    with st.form("quiz_form"):
        for idx, item in enumerate(st.session_state.current_quiz):
            st.write(f"**Q{idx+1}: {item['q']}**")
            user_answers[idx] = st.radio(f"Select option for Q{idx+1}", item["opts"], key=f"q_{idx}", label_visibility="collapsed")
            st.write("")
        submitted = st.form_submit_button("Submit Assessment")
        
    if submitted:
        score = 0
        st.write("---")
        for idx, item in enumerate(st.session_state.current_quiz):
            selected = user_answers[idx].split(")")[0].strip()
            if selected == item["ans"]:
                score += 1
                st.success(f"Q{idx+1}: Correct! {item['exp']}")
            else:
                st.error(f"Q{idx+1}: Incorrect. Correct answer was {item['ans']}. Explanation: {item['exp']}")
        st.metric("Final Score", f"{score} / {len(st.session_state.current_quiz)}")

with tab2:
    st.subheader("Export Student & Class Analytics")
    
    # Excel Download
    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
        students_df.to_excel(writer, index=False, sheet_name='Students')
        lessons_df.to_excel(writer, index=False, sheet_name='Lessons')
    excel_data = excel_buffer.getvalue()
    
    st.download_button(
        label="📥 Download Full Class Summary (.xlsx)",
        data=excel_data,
        file_name="TCS_Class_Summary.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    
    st.markdown("---")
    
    # Word Diagnostic Download
    selected_stu = st.selectbox("Select Student for Diagnostic Dossier", students_df["student_id"] + " - " + students_df["name"])
    stu_record = students_df[students_df["student_id"] == selected_stu.split(" - ")[0]].iloc[0]
    
    if st.button("Generate Official Student Dossier (.docx)"):
        doc = Document()
        doc.add_heading(f"Language Proficiency Diagnostic: {stu_record['name']}", level=1)
        doc.add_paragraph(f"Student ID: {stu_record['student_id']} | Department: {stu_record['department']} | Year: {stu_record['year']}")
        doc.add_paragraph(f"Career Objective: {stu_record['career_goal']} | Overall Rating: {stu_record['level']}")
        
        doc.add_heading("Score Breakdown", level=2)
        for s in ["grammar", "vocabulary", "speaking", "writing", "listening"]:
            doc.add_paragraph(f"{s.capitalize()}: {stu_record[s]}/100")
            
        doc.add_heading("Automated Guidance", level=2)
        doc.add_paragraph("Targeted for weekly study support in lowest scoring competencies.")
        
        docx_buffer = io.BytesIO()
        doc.save(docx_buffer)
        
        st.download_button(
            label="📥 Download Official Word Report (.docx)",
            data=docx_buffer.getvalue(),
            file_name=f"{stu_record['student_id']}_Diagnostic_Report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )